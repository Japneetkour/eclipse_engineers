"""Read a resume (PDF / DOCX / TXT) and turn it into the same features the models use.

    profile = parse_resume(text)  ->  title_family, city, years, skills (the 30 skill_ flags), junior skill scores
Pure rule-based (no internet, no API key). Every guess is shown in the app so the user can correct it.
"""
import io
import re
from datetime import date

# same rules as 01_clean_analytics.py, so a resume lands in the same title family as a posting would
TITLE_RULES = [
    ("Manager", r"manager|director|\bhead\b|chief|vice president|\bvp\b|president|founder|\bcto\b|\bceo\b|owner"),
    ("ML Engineer", r"machine learning|\bml\b|deep learning|\bnlp\b|computer vision|artificial intelligence|\bai\b|chatbot|cognitive"),
    ("Data Engineer", r"data engineer|big data|data architect|hadoop|spark|\betl\b|data platform|database|data warehouse|data pipeline|data integration|informatica|teradata|airflow|nosql|data modell?er"),
    ("BI", r"business intelligence|\bbi\b|tableau|power bi|dashboard|report|visuali[sz]ation|qlik|microstrategy|cognos|\bmis\b"),
    ("Data Scientist", r"data scientist|scientist|data science|statistician|predictive|modell?er|pattern recognition"),
    ("Data Analyst", r"analyst|analytics|business analysis|data analysis"),
]

# words that usually sit in a job title line on a resume
TITLE_WORDS = r"engineer|scientist|analyst|developer|manager|consultant|intern|lead|architect|specialist|associate|director|head|executive|officer"

# resume wording -> the 30 skill columns of analytics_clean.csv
SKILL_PATTERNS = {
    "sql": r"\bsql\b|mysql|postgres|pl/?sql|t-?sql|sql server|sqlite|oracle db",
    "analytics": r"\banalytics\b",
    "python": r"\bpython\b|pandas|numpy|scikit|sklearn|jupyter",
    "finance": r"\bfinance\b|financial services|banking|fintech",
    "java": r"\bjava\b(?!\s*script)|spring boot|\bj2ee\b",
    "r": r"\br\b(?![&\-])(?=[\s,;/)]|$)|rstudio|tidyverse|ggplot|\bshiny\b",
    "sas": r"\bsas\b|sas base|sas enterprise",
    "business_analysis": r"business analys[ie]s|business analyst|requirement gathering|\bbrd\b",
    "machine_learning": r"machine learning|\bml\b|deep learning|tensorflow|pytorch|keras|xgboost|random forest|regression model|classification|neural network|\bnlp\b|computer vision",
    "data_analysis": r"data analysis|analy[sz]ing data|exploratory data|\beda\b",
    "digital_marketing": r"digital marketing|google ads|adwords|social media marketing|\bsem\b|performance marketing",
    "javascript": r"javascript|\bjs\b|node\.?js|react|angular|vue\.?js|typescript|jquery",
    "project_management": r"project management|\bpmp\b|scrum|agile|jira|prince2|project manager",
    "seo": r"\bseo\b|search engine optimi[sz]ation",
    "data_analytics": r"data analytics",
    "sales": r"\bsales\b",
    "excel": r"\bexcel\b|spreadsheet|vlookup|pivot table",
    "accounting": r"accounting|accountant|tally|bookkeeping|\bgaap\b|\bifrs\b",
    "outsourcing": r"outsourcing|\bbpo\b|\bkpo\b",
    "marketing": r"\bmarketing\b",
    "html": r"\bhtml5?\b|\bcss3?\b",
    "hadoop": r"hadoop|\bhdfs\b|\bhive\b|mapreduce|\bpig\b",
    "spark": r"\bspark\b|pyspark|databricks",
    "business_process": r"business process|process improvement|\bbpm\b|six sigma|lean",
    "data_science": r"data science|data scientist",
    "big_data": r"big data|kafka|hbase|cassandra|mongodb|nosql",
    "business_development": r"business development|\bbd\b|lead generation|client acquisition",
    "auditing": r"auditing|\baudit\b|internal audit",
    "financial_analysis": r"financial analys[ie]s|financial model|valuation|\bfp&a\b",
    "market_research": r"market research|consumer insight|survey analysis",
}

# keywords behind the five junior skill scores of the JDS file (each scored 1-5)
JUNIOR_PATTERNS = {
    "big_data_skills": r"hadoop|spark|pyspark|hive|kafka|big data|databricks|aws|azure|gcp|snowflake|airflow|etl|redshift|bigquery|nosql|mongodb",
    "maths_stats_skills": r"statistic|probability|hypothesis|a/b test|regression|linear algebra|calculus|mathemat|econometric|bayesian|time series|forecast|anova|optimi[sz]ation",
    "coding_skills": r"\bpython\b|\bjava\b|c\+\+|\bc#|\bsql\b|\bgit\b|javascript|\bscala\b|\bbash\b|\bgo\b|docker|api|flask|django|leetcode|\br\b",
    "ai_and_ml_skills": r"machine learning|deep learning|tensorflow|pytorch|keras|scikit|xgboost|\bnlp\b|computer vision|neural|llm|transformer|\bgpt|generative|random forest|clustering",
    "dashboard_and_storytelling_skills": r"tableau|power bi|dashboard|looker|qlik|visuali[sz]|storytelling|presentation|stakeholder|matplotlib|seaborn|plotly|report",
}

CITY_PATTERNS = {
    "Bengaluru": r"bengaluru|bangalore", "Mumbai": r"(?<!navi )mumbai|bombay", "Navi Mumbai": r"navi mumbai",
    "Delhi NCR": r"\bdelhi\b|\bncr\b", "Gurgaon": r"gurgaon|gurugram", "Noida": r"noida", "Pune": r"\bpune\b",
    "Hyderabad": r"hyderabad|secunderabad", "Chennai": r"chennai|madras", "Ahmedabad": r"ahmedabad",
    "Kolkata": r"kolkata|calcutta", "Trivandrum": r"trivandrum|thiruvananthapuram", "Kochi": r"kochi|cochin",
    "Chandigarh": r"chandigarh",
}

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


# ---------------------------------------------------------------- text extraction
def extract_text(name, data):
    """name = uploaded file name, data = bytes. Returns plain text."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return "\n".join([p.text for p in d.paragraphs] + cells)
    return data.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------- field guesses
def family_of(title):
    t = str(title).lower()
    for name, pattern in TITLE_RULES:
        if re.search(pattern, t):
            return name
    return "Other"


def guess_title(text):
    """First short line that looks like a job title -> (title line, family)."""
    for line in text.splitlines()[:60]:
        s = line.strip(" •-|:\t")
        if 3 < len(s) < 70 and re.search(TITLE_WORDS, s, re.I) and family_of(s) != "Other":
            return s, family_of(s)
    for line in text.splitlines()[:60]:
        s = line.strip(" •-|:\t")
        if 3 < len(s) < 70 and re.search(TITLE_WORDS, s, re.I):
            return s, family_of(s)
    return "", family_of(text[:1500])


def guess_city(text):
    t = text.lower()
    hits = {c: m.start() for c, p in CITY_PATTERNS.items() if (m := re.search(p, t))}
    return min(hits, key=hits.get) if hits else "Other"  # the earliest mention is usually the contact line


def _ym(mon, yr):
    return int(yr) * 12 + MONTHS.get(str(mon)[:3].lower(), 6) if mon else int(yr) * 12 + 6


def guess_years(text):
    """Stated 'N years of experience' wins; otherwise add up the date ranges of jobs (overlaps merged)."""
    t = text.lower()
    stated = [float(x) for x in re.findall(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs)\b(?:\s+of)?(?:\s+\w+){0,3}\s+experience", t)]
    if stated:
        return round(min(max(stated), 30), 1), "stated in resume"
    now = date.today().year * 12 + date.today().month
    mon = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"
    edu = r"b\.?tech|m\.?tech|\bb\.?e\b|b\.?sc|m\.?sc|bachelor|master|\bmba\b|ph\.?d|degree|university|college|school|institute|cgpa|\bgpa\b|class x"
    work = "\n".join(l for l in t.splitlines() if not re.search(edu, l))  # education dates are not experience
    rng = re.findall(rf"(?:{mon}\s*)?((?:19|20)\d\d)\s*(?:-|–|—|to)\s*(?:{mon}\s*)?((?:19|20)\d\d|present|current|now|till date|date)", work)
    spans = []
    for m1, y1, m2, y2 in rng:
        start = _ym(m1, y1)
        end = now if not y2[:1].isdigit() else _ym(m2, y2)
        if 0 < end - start < 40 * 12:
            spans.append((start, end))
    if not spans:
        return 0.0, "not found (set it yourself)"
    spans.sort()
    total, (cs, ce) = 0, spans[0]
    for s, e in spans[1:]:
        if s <= ce:
            ce = max(ce, e)
        else:
            total += ce - cs; cs, ce = s, e
    total += ce - cs
    return round(total / 12, 1), "summed from job dates - check it"


def guess_skills(text):
    t = text.lower()
    return [s for s, p in SKILL_PATTERNS.items() if re.search(p, t)]


def junior_scores(text):
    """1-5 score per junior skill area on the JDS scale (its scores run 2.2-5): 2.2 = no mention, +0.6 per distinct keyword."""
    t = text.lower()
    out = {}
    for k, p in JUNIOR_PATTERNS.items():
        distinct = len({m.group(0) for m in re.finditer(p, t)})
        out[k] = round(min(5.0, 2.2 + 0.6 * distinct), 1)
    return out


def parse_resume(text):
    title_line, family = guess_title(text)
    years, years_src = guess_years(text)
    return {"title_line": title_line, "title_family": family, "city": guess_city(text),
            "years": years, "years_source": years_src, "skills": guess_skills(text),
            "junior": junior_scores(text), "words": len(text.split())}
